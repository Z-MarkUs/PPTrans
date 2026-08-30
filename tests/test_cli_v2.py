"""Offline end-to-end tests for the PPTrans v2 command-line interface."""

from __future__ import annotations

import json
import socket
from dataclasses import replace
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt
from typer.testing import CliRunner, Result

from pptrans import __version__
from pptrans import cli as cli_module
from pptrans.adapters import providers as providers_module
from pptrans.cli import app

PRIVATE_TEXT = "Confidential launch plan for Project Juniper"


def _create_private_deck(path: Path, *, text: str = PRIVATE_TEXT) -> bytes:
    """Create a minimal, author-owned deck and return its original bytes."""

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(0.75), Inches(0.75), Inches(8), Inches(1.5))
    textbox.name = "Private source text"
    paragraph = textbox.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.runs[0].font.name = "Aptos"
    paragraph.runs[0].font.size = Pt(24)
    paragraph.runs[0].font.bold = True
    presentation.save(path)
    return path.read_bytes()


def _add_unsupported_chart(path: Path) -> bytes:
    presentation = Presentation(path)
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Series", (1, 2))
    presentation.slides[0].shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(3),
        Inches(5),
        Inches(2),
        data,
    )
    presentation.save(path)
    return path.read_bytes()


def _json_output(result: Result) -> object:
    return json.loads(result.stdout)


def _translation_args(
    source: Path,
    output: Path,
    *extra: str,
) -> list[str]:
    return [
        "translate",
        str(source),
        "--source",
        "en",
        "--target",
        "fr",
        "--provider",
        "identity",
        "--output",
        str(output),
        "--json",
        *extra,
    ]


def _dry_run_args(source: Path, *extra: str) -> list[str]:
    return [
        "translate",
        str(source),
        "--source",
        "en",
        "--target",
        "fr",
        "--provider",
        "openai",
        "--model",
        "explicit-preview-model",
        "--dry-run",
        "--json",
        *extra,
    ]


@pytest.fixture
def runner() -> CliRunner:
    return CliRunner()


def test_help_and_version(runner: CliRunner) -> None:
    help_result = runner.invoke(app, ["--help"])

    assert help_result.exit_code == 0, help_result.output
    assert "Translate editable PPTX files" in help_result.stdout
    assert all(command in help_result.stdout for command in ("inspect", "translate", "doctor"))

    translate_help = runner.invoke(app, ["translate", "--help"])
    assert translate_help.exit_code == 0, translate_help.output
    assert "--dry-run" in translate_help.stdout

    version_result = runner.invoke(app, ["--version"])
    assert version_result.exit_code == 0, version_result.output
    assert version_result.stdout.strip() == f"pptrans {__version__}"


def test_inspect_json_is_content_private_by_default_and_opt_in_reveals_text(
    tmp_path: Path,
    runner: CliRunner,
) -> None:
    source = tmp_path / "private.pptx"
    original = _create_private_deck(source)
    base_args = [
        "inspect",
        str(source),
        "--source",
        "en",
        "--target",
        "fr",
        "--json",
    ]

    private_result = runner.invoke(app, base_args)

    assert private_result.exit_code == 0, private_result.output
    private_payload = _json_output(private_result)
    assert isinstance(private_payload, dict)
    assert private_payload["slides"] == 1
    assert private_payload["translation_units"] == 1
    assert private_payload["translatable_spans"] == 1
    assert "units" not in private_payload
    assert PRIVATE_TEXT not in private_result.stdout
    assert source.read_bytes() == original

    strict_private_result = runner.invoke(app, [*base_args, "--fail-on-warnings"])

    assert strict_private_result.exit_code == 0, strict_private_result.output
    assert _json_output(strict_private_result) == private_payload
    assert source.read_bytes() == original

    disclosed_result = runner.invoke(app, [*base_args, "--show-text"])

    assert disclosed_result.exit_code == 0, disclosed_result.output
    disclosed_payload = _json_output(disclosed_result)
    assert isinstance(disclosed_payload, dict)
    assert disclosed_payload["units"][0]["text"] == PRIVATE_TEXT
    assert PRIVATE_TEXT in disclosed_result.stdout
    assert source.read_bytes() == original


def test_inspect_can_fail_on_actionable_unsupported_content_warnings(
    tmp_path: Path,
    runner: CliRunner,
) -> None:
    source = tmp_path / "chart-and-text.pptx"
    _create_private_deck(source)
    original = _add_unsupported_chart(source)
    base_args = [
        "inspect",
        str(source),
        "--source",
        "en",
        "--target",
        "fr",
        "--json",
    ]

    permissive = runner.invoke(app, base_args)

    assert permissive.exit_code == 0, permissive.output
    permissive_payload = _json_output(permissive)
    assert isinstance(permissive_payload, dict)
    warning = permissive_payload["warnings"][0]
    assert warning["code"] == "unsupported_graphic_frame"
    assert warning["slide"] == 1
    assert warning["shape_id_path"]
    assert PRIVATE_TEXT not in permissive.stdout

    strict_json = runner.invoke(app, [*base_args, "--fail-on-warnings"])

    assert strict_json.exit_code == 1
    assert _json_output(strict_json) == permissive_payload
    assert strict_json.stderr == ""

    strict_human = runner.invoke(
        app,
        [arg for arg in base_args if arg != "--json"] + ["--fail-on-warnings"],
    )

    assert strict_human.exit_code == 1
    assert "unsupported_graphic_frame" in strict_human.stderr
    assert "slide 1" in strict_human.stderr
    assert "shape " in strict_human.stderr
    assert PRIVATE_TEXT not in strict_human.stderr
    assert source.read_bytes() == original


def test_strict_translation_stops_before_provider_construction(
    tmp_path: Path,
    runner: CliRunner,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "chart-and-text.pptx"
    output = tmp_path / "translated.pptx"
    _create_private_deck(source)
    original = _add_unsupported_chart(source)
    provider_constructed = False

    def unexpected_provider(*_args: object, **_kwargs: object) -> object:
        nonlocal provider_constructed
        provider_constructed = True
        raise AssertionError("provider must not be constructed")

    monkeypatch.setattr(cli_module, "create_translator", unexpected_provider)

    result = runner.invoke(
        app,
        _translation_args(source, output, "--no-memory", "--fail-on-warnings"),
    )

    assert result.exit_code == 1
    assert provider_constructed is False
    assert "unsupported_graphic_frame" in result.stderr
    assert "--fail-on-warnings is set" in result.stderr
    assert not output.exists()
    assert source.read_bytes() == original


def test_dry_run_reports_a_deck_text_free_upper_bound_without_side_effects(
    tmp_path: Path,
    runner: CliRunner,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "private.pptx"
    original = _create_private_deck(source)

    def unexpected_side_effect(*_args: object, **_kwargs: object) -> object:
        raise AssertionError("dry run touched a side-effecting translation boundary")

    for name in (
        "create_translator",
        "default_memory_path",
        "load_environment",
        "preflight_output",
        "SQLiteTranslationMemory",
        "write_translated_deck",
    ):
        monkeypatch.setattr(cli_module, name, unexpected_side_effect)
    monkeypatch.setattr(providers_module, "import_module", unexpected_side_effect)
    monkeypatch.setattr(socket, "socket", unexpected_side_effect)

    result = runner.invoke(app, _dry_run_args(source))

    assert result.exit_code == 0, result.output
    payload = _json_output(result)
    assert isinstance(payload, dict)
    assert payload["mode"] == "dry-run"
    assert payload["provider"] == "openai"
    assert payload["model"] == "explicit-preview-model"
    assert payload["slides"] == 1
    assert payload["translation_units"] == payload["provider_units"] == 1
    assert payload["translatable_spans"] == 1
    assert payload["provider_calls"] == 1
    assert payload["source_context_characters"] == len(PRIVATE_TEXT)
    assert payload["request_characters"] == payload["largest_request_characters"]
    assert payload["request_characters_per_call"] == [payload["request_characters"]]
    assert payload["assumes_zero_memory_hits"] is True
    assert payload["measurement_scope"] == "characters-not-tokens-cost-or-latency"
    assert payload["warnings"] == []
    assert PRIVATE_TEXT not in result.stdout
    assert str(source) not in result.stdout
    assert source.read_bytes() == original
    assert not source.with_name("private.fr.pptx").exists()


def test_dry_run_human_output_states_its_upper_bound_and_side_effect_boundary(
    tmp_path: Path,
    runner: CliRunner,
) -> None:
    source = tmp_path / "preview.pptx"
    _create_private_deck(source)

    result = runner.invoke(
        app,
        [
            "translate",
            str(source),
            "--source",
            "en",
            "--target",
            "fr",
            "--provider",
            "identity",
            "--dry-run",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "PPTrans provider work preview" in result.stdout
    assert "identity-v1" in result.stdout
    normalized_output = " ".join(result.stdout.split())
    assert "Plan assumes zero translation-memory hits" in normalized_output
    assert "total workload is an upper bound" in normalized_output
    assert "per-call grouping describes this zero-hit plan" in normalized_output
    assert (
        "No credential, provider client, translation memory, output path, or provider/API "
        "network request was used" in normalized_output
    )
    assert PRIVATE_TEXT not in result.stdout


def test_dry_run_redacts_private_glossary_terms_on_validation_failure(
    tmp_path: Path,
    runner: CliRunner,
) -> None:
    source = tmp_path / "preview.pptx"
    original = _create_private_deck(source)
    private_term = "Project Juniper acquisition code name"
    private_target = "Confidential target phrase"
    glossary = tmp_path / "private-glossary.json"
    glossary.write_text(
        json.dumps(
            {
                "terms": [
                    {"source": private_term, "target": private_target},
                    {"source": private_term.upper(), "target": "second target"},
                ]
            }
        ),
        encoding="utf-8",
    )

    result = runner.invoke(app, _dry_run_args(source, "--glossary", str(glossary)))

    assert result.exit_code == 2
    assert "Glossary repeats a source term" in result.stderr
    assert private_term not in result.output
    assert private_term.upper() not in result.output
    assert private_target not in result.output
    assert PRIVATE_TEXT not in result.output
    assert source.read_bytes() == original
    assert not source.with_name("preview.fr.pptx").exists()


def test_dry_run_budget_failure_stops_before_provider_boundaries(
    tmp_path: Path,
    runner: CliRunner,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "preview.pptx"
    original = _create_private_deck(source)

    def unexpected_side_effect_boundary(*_args: object, **_kwargs: object) -> object:
        raise AssertionError("budget failure crossed a side-effect boundary")

    for name in (
        "create_translator",
        "default_memory_path",
        "load_environment",
        "preflight_output",
        "SQLiteTranslationMemory",
        "write_translated_deck",
    ):
        monkeypatch.setattr(cli_module, name, unexpected_side_effect_boundary)
    monkeypatch.setattr(providers_module, "import_module", unexpected_side_effect_boundary)
    monkeypatch.setattr(socket, "socket", unexpected_side_effect_boundary)

    result = runner.invoke(
        app,
        _dry_run_args(source, "--max-provider-request-characters", "1"),
    )

    assert result.exit_code == 2
    assert "request characters; limit is 1" in result.stderr
    assert PRIVATE_TEXT not in result.output
    assert source.read_bytes() == original
    assert not source.with_name("preview.fr.pptx").exists()


@pytest.mark.parametrize(
    ("arguments", "message"),
    [
        (("--output", "preview.pptx"), "--output or --overwrite"),
        (("--overwrite",), "--output or --overwrite"),
        (("--env-file", "credentials.env"), "--env-file"),
        (("--memory", "memory.sqlite"), "--memory"),
    ],
    ids=["output", "overwrite", "environment", "memory"],
)
def test_dry_run_rejects_options_that_imply_side_effects(
    tmp_path: Path,
    runner: CliRunner,
    arguments: tuple[str, ...],
    message: str,
) -> None:
    source = tmp_path / "preview.pptx"
    _create_private_deck(source)

    result = runner.invoke(app, _dry_run_args(source, *arguments))

    assert result.exit_code == 2
    assert message in result.stderr
    assert source.read_bytes()
    assert not (tmp_path / "preview.fr.pptx").exists()
    assert not (tmp_path / "memory.sqlite").exists()


def test_strict_dry_run_stops_on_inspection_warnings_before_estimation(
    tmp_path: Path,
    runner: CliRunner,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "chart-and-text.pptx"
    _create_private_deck(source)
    original = _add_unsupported_chart(source)

    def unexpected_estimate(*_args: object, **_kwargs: object) -> object:
        raise AssertionError("strict warning gate must stop before estimation")

    monkeypatch.setattr(cli_module, "estimate_provider_work", unexpected_estimate)

    result = runner.invoke(app, _dry_run_args(source, "--fail-on-warnings"))

    assert result.exit_code == 1
    assert "unsupported_graphic_frame" in result.stderr
    assert "--fail-on-warnings is set" in result.stderr
    assert PRIVATE_TEXT not in result.stderr
    assert source.read_bytes() == original


def test_identity_translation_is_offline_end_to_end_and_preserves_source(
    tmp_path: Path,
    runner: CliRunner,
) -> None:
    source = tmp_path / "source.pptx"
    output = tmp_path / "translated.pptx"
    original = _create_private_deck(source)

    result = runner.invoke(app, _translation_args(source, output, "--no-memory"))

    assert result.exit_code == 0, result.output
    payload = _json_output(result)
    assert isinstance(payload, dict)
    assert payload["output"] == str(output.resolve())
    assert payload["provider"] == "identity"
    assert payload["model"] == "identity-v1"
    assert payload["memory_hits"] == 0
    assert payload["provider_calls"] == 1
    assert payload["verified_spans"] == 1
    assert source.read_bytes() == original
    assert output.exists()

    reopened = Presentation(output)
    paragraph = reopened.slides[0].shapes[0].text_frame.paragraphs[0]
    assert paragraph.text == PRIVATE_TEXT
    assert paragraph.runs[0].font.bold is True
    assert paragraph.runs[0].font.size.pt == 24


def test_translation_refuses_to_overwrite_its_source(
    tmp_path: Path,
    runner: CliRunner,
) -> None:
    source = tmp_path / "source.pptx"
    original = _create_private_deck(source)

    result = runner.invoke(app, _translation_args(source, source, "--no-memory", "--overwrite"))

    assert result.exit_code == 2
    assert "must not overwrite the source" in result.output
    assert source.read_bytes() == original
    assert Presentation(source).slides[0].shapes[0].text == PRIVATE_TEXT


def test_existing_output_requires_opt_in_and_overwrite_then_replaces_it(
    tmp_path: Path,
    runner: CliRunner,
) -> None:
    source = tmp_path / "source.pptx"
    output = tmp_path / "existing.pptx"
    _create_private_deck(source)
    output.write_bytes(b"sentinel: do not replace")

    refused = runner.invoke(app, _translation_args(source, output, "--no-memory"))

    assert refused.exit_code == 2
    assert "Output already exists" in refused.output
    assert output.read_bytes() == b"sentinel: do not replace"

    replaced = runner.invoke(
        app,
        _translation_args(source, output, "--no-memory", "--overwrite"),
    )

    assert replaced.exit_code == 0, replaced.output
    assert output.read_bytes() != b"sentinel: do not replace"
    assert Presentation(output).slides[0].shapes[0].text == PRIVATE_TEXT


def test_destination_preflight_fails_before_provider_construction(
    tmp_path: Path,
    runner: CliRunner,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "source.pptx"
    output = tmp_path / "existing.pptx"
    _create_private_deck(source)
    output.write_bytes(b"preserve me")

    monkeypatch.setattr(
        cli_module,
        "create_translator",
        lambda *_args, **_kwargs: pytest.fail("provider must not be constructed"),
    )

    result = runner.invoke(app, _translation_args(source, output, "--no-memory"))

    assert result.exit_code == 2
    assert "Output already exists" in result.output
    assert output.read_bytes() == b"preserve me"


@pytest.mark.parametrize("overwrite", [False, True], ids=["no-overwrite", "overwrite"])
def test_cli_rejects_output_symlink_without_touching_its_target(
    tmp_path: Path,
    runner: CliRunner,
    monkeypatch: pytest.MonkeyPatch,
    *,
    overwrite: bool,
) -> None:
    source = tmp_path / "source.pptx"
    output = tmp_path / "redirected.pptx"
    sentinel = tmp_path / "sentinel.bin"
    _create_private_deck(source)
    sentinel.write_bytes(b"do not replace")
    try:
        output.symlink_to(sentinel)
    except OSError as exc:
        pytest.skip(f"symbolic links unavailable on this filesystem: {exc}")
    monkeypatch.setattr(
        cli_module,
        "create_translator",
        lambda *_args, **_kwargs: pytest.fail("provider must not be constructed"),
    )
    args = _translation_args(source, output, "--no-memory")
    if overwrite:
        args.append("--overwrite")

    result = runner.invoke(app, args)

    assert result.exit_code == 2
    assert "symbolic link" in result.output
    assert sentinel.read_bytes() == b"do not replace"
    assert output.is_symlink()


def test_provider_budget_fails_before_provider_construction(
    tmp_path: Path,
    runner: CliRunner,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "source.pptx"
    output = tmp_path / "output.pptx"
    _create_private_deck(source)
    monkeypatch.setattr(
        cli_module,
        "create_translator",
        lambda *_args, **_kwargs: pytest.fail("provider must not be constructed"),
    )

    result = runner.invoke(
        app,
        _translation_args(
            source,
            output,
            "--no-memory",
            "--max-provider-source-characters",
            "1",
        ),
    )

    assert result.exit_code == 2
    assert "source/context characters" in result.output
    assert not output.exists()


def test_schema_impossible_span_count_fails_before_provider_construction(
    tmp_path: Path,
    runner: CliRunner,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "source.pptx"
    output = tmp_path / "output.pptx"
    _create_private_deck(source)
    real_plan = cli_module.inspect_deck(source, source_lang="en", target_lang="fr")
    unit = real_plan.units[0]
    template = unit.spans[0]
    oversized = replace(
        real_plan,
        units=(
            replace(
                unit,
                spans=tuple(
                    replace(template, id=f"oversized-{index}", node_index=index)
                    for index in range(10_001)
                ),
            ),
        ),
    )
    monkeypatch.setattr(cli_module, "inspect_deck", lambda *_args, **_kwargs: oversized)
    monkeypatch.setattr(
        cli_module,
        "create_translator",
        lambda *_args, **_kwargs: pytest.fail("provider must not be constructed"),
    )

    result = runner.invoke(app, _translation_args(source, output, "--no-memory"))

    assert result.exit_code == 2
    assert "too many translatable spans" in result.output
    assert not output.exists()


def test_explicit_memory_reuses_translation_and_no_memory_creates_no_database(
    tmp_path: Path,
    runner: CliRunner,
) -> None:
    source = tmp_path / "source.pptx"
    memory = tmp_path / "cache" / "translations.sqlite3"
    disabled_memory = tmp_path / "disabled" / "translations.sqlite3"
    _create_private_deck(source, text="Cache this sentence")

    first = runner.invoke(
        app,
        _translation_args(
            source,
            tmp_path / "first.pptx",
            "--memory",
            str(memory),
        ),
    )
    second = runner.invoke(
        app,
        _translation_args(
            source,
            tmp_path / "second.pptx",
            "--memory",
            str(memory),
        ),
    )
    uncached = runner.invoke(
        app,
        _translation_args(
            source,
            tmp_path / "uncached.pptx",
            "--no-memory",
        ),
    )

    assert first.exit_code == 0, first.output
    assert second.exit_code == 0, second.output
    assert uncached.exit_code == 0, uncached.output
    first_payload = _json_output(first)
    second_payload = _json_output(second)
    uncached_payload = _json_output(uncached)
    assert isinstance(first_payload, dict)
    assert isinstance(second_payload, dict)
    assert isinstance(uncached_payload, dict)
    assert first_payload["memory_hits"] == 0
    assert first_payload["provider_calls"] == 1
    assert second_payload["memory_hits"] == 1
    assert second_payload["provider_calls"] == 0
    assert uncached_payload["memory_hits"] == 0
    assert uncached_payload["provider_calls"] == 1
    assert memory.is_file()
    assert not disabled_memory.exists()


@pytest.mark.parametrize("collision", ["source", "output"])
def test_translation_memory_cannot_alias_a_deck_path(
    tmp_path: Path,
    runner: CliRunner,
    monkeypatch: pytest.MonkeyPatch,
    collision: str,
) -> None:
    source = tmp_path / "source.pptx"
    output = tmp_path / "output.pptx"
    _create_private_deck(source)
    memory = source if collision == "source" else output
    monkeypatch.setattr(
        cli_module,
        "create_translator",
        lambda *_args, **_kwargs: pytest.fail("provider must not be constructed"),
    )

    result = runner.invoke(
        app,
        _translation_args(source, output, "--memory", str(memory)),
    )

    assert result.exit_code == 2
    assert "Translation-memory path must be distinct" in result.output
    assert not output.exists()


def test_memory_and_no_memory_are_mutually_exclusive(
    tmp_path: Path,
    runner: CliRunner,
) -> None:
    source = tmp_path / "source.pptx"
    output = tmp_path / "output.pptx"
    _create_private_deck(source)

    result = runner.invoke(
        app,
        _translation_args(
            source,
            output,
            "--memory",
            str(tmp_path / "cache.sqlite3"),
            "--no-memory",
        ),
    )

    assert result.exit_code == 2
    assert "cannot be combined" in result.output
    assert not output.exists()


def test_cli_rejects_symlinked_memory_before_provider_construction(
    tmp_path: Path,
    runner: CliRunner,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "source.pptx"
    output = tmp_path / "output.pptx"
    target = tmp_path / "target.sqlite3"
    memory = tmp_path / "selected.sqlite3"
    _create_private_deck(source)
    target.write_bytes(b"preserve target")
    try:
        memory.symlink_to(target)
    except OSError as exc:
        pytest.skip(f"symbolic links unavailable on this filesystem: {exc}")
    monkeypatch.setattr(
        cli_module,
        "create_translator",
        lambda *_args, **_kwargs: pytest.fail("provider must not be constructed"),
    )

    result = runner.invoke(
        app,
        _translation_args(source, output, "--memory", str(memory)),
    )

    assert result.exit_code == 2
    assert "must not be a symbolic link" in result.output
    assert target.read_bytes() == b"preserve target"
    assert not output.exists()


@pytest.mark.parametrize("kind", ["environment", "glossary"])
def test_output_cannot_overwrite_configuration_inputs(
    tmp_path: Path,
    runner: CliRunner,
    kind: str,
) -> None:
    source = tmp_path / "source.pptx"
    protected = tmp_path / f"protected-{kind}.pptx"
    _create_private_deck(source)
    original = b"OPENAI_API_KEY=opaque\n" if kind == "environment" else b'{"terms": []}\n'
    protected.write_bytes(original)
    option = "--env-file" if kind == "environment" else "--glossary"

    result = runner.invoke(
        app,
        _translation_args(
            source,
            protected,
            option,
            str(protected),
            "--overwrite",
            "--no-memory",
        ),
    )

    assert result.exit_code == 2
    assert "must not overwrite an environment or glossary input" in result.output
    assert protected.read_bytes() == original


def test_unusable_memory_path_is_reported_as_an_expected_cli_error(
    tmp_path: Path,
    runner: CliRunner,
) -> None:
    source = tmp_path / "source.pptx"
    output = tmp_path / "output.pptx"
    unusable_memory = tmp_path / "memory-is-a-directory"
    _create_private_deck(source)
    unusable_memory.mkdir()

    result = runner.invoke(
        app,
        _translation_args(
            source,
            output,
            "--memory",
            str(unusable_memory),
        ),
    )

    assert result.exit_code == 2
    assert "Error:" in result.output
    assert "translation-memory" in result.output.lower()
    assert not output.exists()


@pytest.mark.parametrize(
    ("provider", "environment_name"),
    [("openai", "OPENAI_API_KEY"), ("anthropic", "ANTHROPIC_API_KEY")],
)
def test_paid_provider_credential_failure_is_explicit_and_offline(
    tmp_path: Path,
    runner: CliRunner,
    monkeypatch: pytest.MonkeyPatch,
    provider: str,
    environment_name: str,
) -> None:
    source = tmp_path / "source.pptx"
    output = tmp_path / "should-not-exist.pptx"
    env_file = tmp_path / "empty.env"
    _create_private_deck(source)
    env_file.write_text("# intentionally empty\n", encoding="utf-8")
    monkeypatch.delenv(environment_name, raising=False)

    result = runner.invoke(
        app,
        [
            "translate",
            str(source),
            "--source",
            "en",
            "--target",
            "fr",
            "--provider",
            provider,
            "--model",
            "offline-contract-test",
            "--env-file",
            str(env_file),
            "--output",
            str(output),
            "--no-memory",
        ],
    )

    assert result.exit_code == 2
    assert environment_name in result.output
    assert "Error:" in result.output
    assert not output.exists()
