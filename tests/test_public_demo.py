"""Keep the committed public demo synchronized with the real v2 pipeline."""

from __future__ import annotations

import hashlib
import importlib.util
import json
import sys
from pathlib import Path
from types import ModuleType
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation

from pptrans.adapters.providers.identity import IdentityTranslator
from pptrans.application.deck import write_translated_deck
from pptrans.application.translate import (
    TranslationOptions,
    estimate_provider_work,
    translate_plan,
)
from pptrans.ooxml import inspect_deck

REPO_ROOT = Path(__file__).parents[1]
DEMO_PATH = REPO_ROOT / "examples" / "pptrans-demo.en.pptx"
CURATED_DEMO_PATH = REPO_ROOT / "examples" / "pptrans-demo.zh-CN.pptx"
SOURCE_MANIFEST_PATH = REPO_ROOT / "examples" / "pptrans-demo.source" / "manifest.json"
QA_PATH = REPO_ROOT / "docs" / "qa" / "2026-08-31-exact-rebuild.json"
HISTORICAL_NATIVE_QA_PATH = REPO_ROOT / "docs" / "qa" / "2026-08-28-windows-libreoffice.json"
HISTORICAL_CURATED_QA_PATH = REPO_ROOT / "docs" / "qa" / "2026-08-28-curated-zh-cn.json"
QA_CANONICAL_SHA256 = "462eef3c18c1cd00dfbf4e7a6521791564c4962fb7fc88bdb84eeff7a436270b"
PUBLIC_REPOSITORY_URL = "https://github.com/Z-MarkUs/PPTrans"
SOURCE_SHOWCASE_STATUS = "PPTrans v2 • unreleased local showcase"
TARGET_SHOWCASE_STATUS = "PPTrans v2 • 未发布的本地展示版"


def _canonical_record_sha256(record: object) -> str:
    payload = json.dumps(record, ensure_ascii=True, sort_keys=True, separators=(",", ":")).encode()
    return hashlib.sha256(payload).hexdigest()


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _load_curated_demo_script() -> ModuleType:
    path = REPO_ROOT / "scripts" / "build_curated_demo.py"
    spec = importlib.util.spec_from_file_location("test_build_curated_demo", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def test_public_demo_runs_end_to_end_without_network(tmp_path: Path) -> None:
    plan = inspect_deck(DEMO_PATH, source_lang="en", target_lang="en")

    assert len(plan.slide_parts) == 3
    assert len(plan.units) == 41
    assert sum(len(unit.translatable_span_ids) for unit in plan.units) == 45
    assert plan.warnings == ()

    run = translate_plan(plan, IdentityTranslator())
    result = write_translated_deck(plan, run.translations, tmp_path / "demo.identity.pptx")

    assert result.report.verified_patches == 41
    assert result.report.verified_spans == 45
    assert run.stats.provider_calls == 2
    assert run.stats.input_tokens == 0
    assert result.output_path.read_bytes() == DEMO_PATH.read_bytes()
    reopened = Presentation(str(result.output_path))
    assert len(reopened.slides) == 3
    assert reopened.core_properties.author == "Hehan Zhao"
    assert reopened.core_properties.last_modified_by == "Hehan Zhao"
    assert reopened.core_properties.title == "PPTrans v2 - verifiable OOXML translation demo"
    assert reopened.core_properties.subject.startswith("Synthetic fixture")


def test_pre_release_demo_status_and_link_boundary() -> None:
    """Keep the dated recruiter visual scoped as a pre-release showcase."""

    for deck_path, expected_status in (
        (DEMO_PATH, SOURCE_SHOWCASE_STATUS),
        (CURATED_DEMO_PATH, TARGET_SHOWCASE_STATUS),
    ):
        with ZipFile(deck_path) as archive:
            slide = archive.read("ppt/slides/slide1.xml").decode("utf-8-sig")
            relationships = archive.read("ppt/slides/_rels/slide1.xml.rels").decode("utf-8-sig")
        assert expected_status in slide
        assert PUBLIC_REPOSITORY_URL not in slide
        assert PUBLIC_REPOSITORY_URL not in relationships
        assert "relationships/hyperlink" not in relationships

    readme = (REPO_ROOT / "README.md").read_text(encoding="utf-8")
    assert "The cover identifies v2 as a pre-release showcase" in readme


def test_public_demo_provider_preview_is_exact_and_deck_text_free() -> None:
    plan = inspect_deck(DEMO_PATH, source_lang="en", target_lang="zh-CN")

    estimate = estimate_provider_work(
        plan.units,
        TranslationOptions(),
        source_lang=plan.source_lang,
        target_lang=plan.target_lang,
    )

    assert (
        estimate.provider_units,
        estimate.provider_calls,
        estimate.source_context_characters,
        estimate.request_characters,
        estimate.largest_request_characters,
        estimate.request_characters_per_call,
    ) == (41, 2, 2_799, 8_867, 5_903, (5_903, 2_964))


def test_current_qa_record_is_pinned_to_the_exact_rebuild_and_native_assets() -> None:
    record = json.loads(QA_PATH.read_bytes())
    manifest = json.loads(SOURCE_MANIFEST_PATH.read_bytes())
    source_bytes = DEMO_PATH.read_bytes()
    target_bytes = CURATED_DEMO_PATH.read_bytes()

    assert _canonical_record_sha256(record) == QA_CANONICAL_SHA256
    assert record["schema_version"] == 1
    assert record["exact_rebuild"] == {
        "command": "python scripts/rebuild_demo.py --check",
        "builder_path": "scripts/rebuild_demo.py",
        "builder_bytes": (REPO_ROOT / "scripts" / "rebuild_demo.py").stat().st_size,
        "builder_sha256": _sha256(REPO_ROOT / "scripts" / "rebuild_demo.py"),
        "source_directory": "examples/pptrans-demo.source",
        "manifest_path": "examples/pptrans-demo.source/manifest.json",
        "manifest_bytes": SOURCE_MANIFEST_PATH.stat().st_size,
        "manifest_sha256": _sha256(SOURCE_MANIFEST_PATH),
        "entry_count": 29,
        "archive_compression": "ZIP_STORED",
        "fixed_timestamp": "2026-08-28T00:34:00",
        "standard_library_only": True,
        "byte_identical_across_two_local_rebuilds": True,
        "local_windows_python_matrix": [
            {
                "python": "3.10.11",
                "distribution": "python-3.10.11-embed-amd64.zip",
                "distribution_url": (
                    "https://www.python.org/ftp/python/3.10.11/python-3.10.11-embed-amd64.zip"
                ),
                "distribution_sha256": (
                    "608619f8619075629c9c69f361352a0da6ed7e62f83a0e19c63e0ea32eb7629d"
                ),
                "result_bytes": len(source_bytes),
                "result_sha256": hashlib.sha256(source_bytes).hexdigest(),
                "byte_identical": True,
            },
            {
                "python": "3.12.13",
                "distribution": "workspace interpreter",
                "result_bytes": len(source_bytes),
                "result_sha256": hashlib.sha256(source_bytes).hexdigest(),
                "byte_identical": True,
            },
            {
                "python": "3.13.15",
                "distribution": "python-3.13.15-embed-amd64.zip",
                "distribution_url": (
                    "https://www.python.org/ftp/python/3.13.15/python-3.13.15-embed-amd64.zip"
                ),
                "distribution_sha256": (
                    "d1f04d990aee1253d8569e8e5104e30fa9f5fa830899f14843448872d936a2cf"
                ),
                "result_bytes": len(source_bytes),
                "result_sha256": hashlib.sha256(source_bytes).hexdigest(),
                "byte_identical": True,
            },
        ],
    }
    assert manifest["expected_package"] == {
        "path": "examples/pptrans-demo.en.pptx",
        "bytes": len(source_bytes),
        "sha256": hashlib.sha256(source_bytes).hexdigest(),
    }
    assert record["presentations"] == {
        "source": {
            "path": "examples/pptrans-demo.en.pptx",
            "bytes": len(source_bytes),
            "sha256": hashlib.sha256(source_bytes).hexdigest(),
        },
        "identity_output": {
            "bytes": len(source_bytes),
            "sha256": hashlib.sha256(source_bytes).hexdigest(),
            "byte_identical_to_source": True,
        },
        "curated_target": {
            "path": "examples/pptrans-demo.zh-CN.pptx",
            "bytes": len(target_bytes),
            "sha256": hashlib.sha256(target_bytes).hexdigest(),
            "byte_identical_to_source": False,
        },
    }
    assert all(record["package_verification"].values())
    assert record["pipeline"]["identity"]["verified_patches"] == 41
    assert record["pipeline"]["identity"]["verified_spans"] == 45
    curated_pipeline = record["pipeline"]["curated_target"]
    curated_builder = REPO_ROOT / "scripts" / "build_curated_demo.py"
    assert curated_pipeline["builder_path"] == "scripts/build_curated_demo.py"
    assert curated_pipeline["builder_bytes"] == curated_builder.stat().st_size
    assert curated_pipeline["builder_sha256"] == _sha256(curated_builder)
    assert curated_pipeline["changed_parts"] == [
        "ppt/slides/slide1.xml",
        "ppt/slides/slide2.xml",
        "ppt/slides/slide3.xml",
    ]
    assert record["renderer"] == {
        "backend": "libreoffice",
        "libreoffice_version": "26.8.0.3",
        "libreoffice_build": "bce0998afefdbc355585ca324285661a2170ba77",
        "distribution": "LibreOffice_26.8.0.3_Win_x86-64.msi",
        "distribution_sha256": "4aa6c6e1895f4055104effcb556bd3362d20c6ad707c149543304f395ef9db95",
        "installation_mode": (
            "administrative extraction into a disposable audit directory; no system installation"
        ),
        "pymupdf_version": "1.28.2",
        "dpi": 144,
        "render_api": (
            "LibreOfficeRenderer(executable=soffice).render(RenderRequest(deck, output_dir, "
            "dpi=144, timeout_seconds=180))"
        ),
        "replay_script_path": "scripts/reproduce_native_demo.py",
        "replay_script_bytes": (REPO_ROOT / "scripts" / "reproduce_native_demo.py").stat().st_size,
        "replay_script_sha256": _sha256(REPO_ROOT / "scripts" / "reproduce_native_demo.py"),
    }

    assert [slide["slide_number"] for slide in record["slides"]] == [1, 2, 3]
    for slide in record["slides"]:
        number = slide["slide_number"]
        assert slide["identity_pixel_equal"] is True
        assert slide["target_dimensions_equal"] is True
        assert slide["source_png_sha256"] == slide["identity_png_sha256"]
        for language, key in (("en", "source"), ("zh-CN", "target")):
            image_path = (
                REPO_ROOT
                / "docs"
                / "assets"
                / f"pptrans-demo-libreoffice-{language}-slide-{number:02d}.png"
            )
            assert _sha256(image_path) == slide[f"{key}_png_sha256"]
            with Image.open(image_path) as image:
                assert image.format == "PNG"
                assert image.size == (slide["width"], slide["height"])

    historical_native = json.loads(HISTORICAL_NATIVE_QA_PATH.read_bytes())
    historical_curated = json.loads(HISTORICAL_CURATED_QA_PATH.read_bytes())
    assert record["lineage"]["prior_source"] == historical_native["presentations"]["source"]
    assert (
        record["lineage"]["prior_curated_target"]
        == historical_curated["presentations"]["curated_target"]
    )


def test_curated_zh_cn_demo_is_reproducible_and_structurally_verified(tmp_path: Path) -> None:
    curated = _load_curated_demo_script()
    rebuilt = tmp_path / "pptrans-demo.zh-CN.pptx"
    result, run = curated.build_curated_demo(DEMO_PATH, rebuilt)

    assert rebuilt.read_bytes() == CURATED_DEMO_PATH.read_bytes()
    assert result.report.changed_parts == (
        "ppt/slides/slide1.xml",
        "ppt/slides/slide2.xml",
        "ppt/slides/slide3.xml",
    )
    assert result.report.verified_patches == 41
    assert result.report.verified_spans == 45
    assert run.stats.provider_calls == 2
    assert run.stats.input_tokens == 0
    assert run.stats.output_tokens == 0

    with ZipFile(DEMO_PATH) as source, ZipFile(rebuilt) as target:
        assert source.namelist() == target.namelist()
        changed_members = {
            name for name in source.namelist() if source.read(name) != target.read(name)
        }
    assert changed_members == set(result.report.changed_parts)

    source_plan = inspect_deck(DEMO_PATH, source_lang="en", target_lang="zh-CN")
    target_plan = inspect_deck(CURATED_DEMO_PATH, source_lang="zh-CN", target_lang="en")
    for source_unit, target_unit in zip(source_plan.units, target_plan.units, strict=True):
        assert target_unit.locator == source_unit.locator
        assert tuple(span.id for span in target_unit.spans) == tuple(
            span.id for span in source_unit.spans
        )
        for source_span, target_span in zip(source_unit.spans, target_unit.spans, strict=True):
            assert target_span.source == curated.CURATED_TRANSLATIONS[source_span.source]

    reopened = Presentation(str(CURATED_DEMO_PATH))
    assert len(reopened.slides) == 3
    assert reopened.core_properties.author == "Hehan Zhao"
