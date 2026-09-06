"""Focused branch coverage for PPTrans command-line presentation paths."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches
from rich.console import Console
from rich.text import Text
from typer.testing import CliRunner

from pptrans import cli
from pptrans.doctor import DoctorCheck


class _ReconfigurableStream(io.StringIO):
    def __init__(self) -> None:
        super().__init__()
        self.configuration: tuple[str, str] | None = None

    def reconfigure(self, *, encoding: str, errors: str) -> None:
        self.configuration = (encoding, errors)


class _TtyCapture(io.StringIO):
    def isatty(self) -> bool:
        return True


def _create_deck(path: Path, text: str = "Human-readable output") -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    textbox.text = text
    presentation.save(path)


def test_configure_utf8_streams_handles_reconfigurable_and_capture_streams(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    reconfigurable = _ReconfigurableStream()
    capture_only = io.StringIO()
    monkeypatch.setattr(cli.sys, "stdout", reconfigurable)
    monkeypatch.setattr(cli.sys, "stderr", capture_only)

    cli._configure_utf8_streams()

    assert reconfigurable.configuration == ("utf-8", "replace")


@pytest.mark.parametrize(
    ("language", "expected_name"),
    [
        (" zh / CN! ", "deck.zh---CN.pptx"),
        (" / ", "deck.translated.pptx"),
        ("pt_BR", "deck.pt_BR.pptx"),
    ],
)
def test_default_output_sanitizes_language_for_a_sibling_path(
    language: str,
    expected_name: str,
) -> None:
    source = Path("folder") / "deck.pptx"

    assert cli._default_output(source, language) == source.with_name(expected_name)


def test_human_inspect_output_includes_summary_and_opted_in_text(tmp_path: Path) -> None:
    source = tmp_path / "deck.pptx"
    _create_deck(source, "Visible only by explicit opt-in")

    result = CliRunner().invoke(
        cli.app,
        [
            "inspect",
            str(source),
            "--source",
            "en",
            "--target",
            "fr",
            "--show-text",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "PPTrans deck plan" in result.stdout
    assert "Translation units" in result.stdout
    assert "Warnings" in result.stdout
    assert "Visible only by explicit opt-in" in result.stdout


def test_inspect_reports_expected_validation_errors(tmp_path: Path) -> None:
    result = CliRunner().invoke(
        cli.app,
        [
            "inspect",
            str(tmp_path / "missing.pptx"),
            "--source",
            "en",
            "--target",
            "fr",
        ],
    )

    assert result.exit_code == 2
    assert "Error:" in result.output


def test_untrusted_paths_and_deck_text_are_rendered_as_literal_text(tmp_path: Path) -> None:
    source = tmp_path / "[bold red]deck.pptx"
    _create_deck(
        source,
        "[link=https://untrusted.invalid]literal[/link]\x9d0;owned\x9c\x1b[31m",
    )

    result = CliRunner().invoke(
        cli.app,
        [
            "inspect",
            str(source),
            "--source",
            "en",
            "--target",
            "fr",
            "--show-text",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "[link=https://untrusted.invalid]literal[/link]" in result.stdout
    assert "\x1b" not in result.stdout
    assert "\x9d" not in result.stdout
    assert "\x9c" not in result.stdout
    assert "\\u009d0;owned\\u009c" in result.stdout


def test_terminal_sanitizer_blocks_ansi_and_osc_even_with_forced_color() -> None:
    output = io.StringIO()
    forced = Console(file=output, force_terminal=True, color_system="standard")

    forced.print(Text(cli._terminal_string("A\x1b[31mB\x1b]52;c;payload\x07")))

    rendered = output.getvalue()
    assert "\x1b[31m" not in rendered
    assert "\x1b]52" not in rendered
    assert "\x07" not in rendered
    assert "\\u001b[31m" in rendered
    assert "\\u001b]52;c;payload\\u0007" in rendered


def test_json_writer_is_plain_valid_json_on_a_tty(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    output = _TtyCapture()
    monkeypatch.setattr(cli.sys, "stdout", output)

    cli._write_json({"text": "before\x9b31mafter", "unicode": "你好"})

    rendered = output.getvalue()
    assert "\x1b" not in rendered
    assert "\x9b" not in rendered
    assert json.loads(rendered) == {"text": "before\x9b31mafter", "unicode": "你好"}


def test_untrusted_error_path_does_not_execute_rich_markup(tmp_path: Path) -> None:
    missing = tmp_path / "[bold]missing.pptx"

    result = CliRunner().invoke(
        cli.app,
        ["inspect", str(missing), "--source", "en", "--target", "fr"],
    )

    assert result.exit_code == 2
    rendered = result.output
    assert "\x1b" not in rendered
    assert "\x9b" not in rendered
    assert "[bold]missing.pptx" in "".join(rendered.split())


def test_human_translation_uses_sanitized_default_output_and_glossary(tmp_path: Path) -> None:
    source = tmp_path / "showcase.pptx"
    glossary = tmp_path / "terms.json"
    _create_deck(source, "gross margin")
    glossary.write_text(
        '{"terms":[{"source":"gross margin","target":"marge brute"}]}',
        encoding="utf-8",
    )

    result = CliRunner().invoke(
        cli.app,
        [
            "translate",
            str(source),
            "--source",
            "en",
            "--target",
            "fr / CA",
            "--provider",
            "identity",
            "--glossary",
            str(glossary),
            "--no-memory",
        ],
    )

    expected_output = tmp_path / "showcase.fr---CA.pptx"
    assert result.exit_code == 0, result.output
    assert "Translated and verified:" in result.stdout
    assert "text spans" in result.stdout
    assert expected_output.is_file()


def test_doctor_human_table_and_required_failure(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    checks = (
        DoctorCheck(name="python", status="pass", detail="3.12", required=True),
        DoctorCheck(name="renderer", status="warn", detail="optional", required=False),
        DoctorCheck(name="credential", status="fail", detail="not set", required=True),
    )
    monkeypatch.setattr(cli, "load_environment", lambda _path: False)
    monkeypatch.setattr(cli, "run_doctor", lambda _provider: checks)

    result = CliRunner().invoke(cli.app, ["doctor", "--provider", "anthropic"])

    assert result.exit_code == 1
    assert "PPTrans doctor" in result.stdout
    assert "PYTHON" not in result.stdout
    assert "python" in result.stdout
    assert "WARN" in result.stdout
    assert "FAIL" in result.stdout


def test_doctor_reports_environment_loading_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail(_path: Path | None) -> bool:
        raise OSError("unreadable environment")

    monkeypatch.setattr(cli, "load_environment", fail)

    result = CliRunner().invoke(cli.app, ["doctor"])

    assert result.exit_code == 2
    assert "unreadable environment" in result.output


def test_main_configures_streams_before_dispatch(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[str] = []
    monkeypatch.setattr(cli, "_configure_utf8_streams", lambda: calls.append("streams"))
    monkeypatch.setattr(cli, "app", lambda: calls.append("app"))

    cli.main()

    assert calls == ["streams", "app"]
