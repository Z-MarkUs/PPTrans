"""Fixtures shared by the v2 OOXML core tests."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Inches, Pt


def create_complex_deck(path: Path) -> None:  # noqa: PLR0915
    """Create an author-owned deck covering the formatting invariants."""

    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    slide = presentation.slides[0]

    textbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5.0), Inches(1.5))
    textbox.name = "Mixed runs"
    frame = textbox.text_frame
    frame.margin_left = Inches(0.13)
    frame.margin_right = Inches(0.17)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.level = 1
    paragraph.space_before = Pt(5)
    paragraph.space_after = Pt(7)
    paragraph.line_spacing = 1.2

    first = paragraph.add_run()
    first.text = "Revenue "
    first.font.name = "Aptos"
    first.font.size = Pt(24)
    first.font.bold = True
    first.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    first.hyperlink.address = "https://example.com/revenue"

    second = paragraph.add_run()
    second.text = "grew 20%"
    second.font.name = "Courier New"
    second.font.size = Pt(18)
    second.font.italic = True
    second.font.color.rgb = RGBColor(0x66, 0x99, 0x22)
    paragraph.add_line_break()
    after_break = paragraph.add_run()
    after_break.text = "Year over year"
    after_break.font.underline = True

    field = parse_xml(
        '<a:fld id="{AF4AF76E-4C38-47F9-ABE3-0B2119F9A266}" type="datetime1" '
        f'{nsdecls("a")}><a:rPr lang="en-US"/><a:t>2026-08-28</a:t></a:fld>'
    )
    paragraph._p.insert(len(paragraph._p) - 1, field)

    second_paragraph = frame.add_paragraph()
    second_paragraph.text = "Second paragraph"
    second_paragraph.level = 2
    second_paragraph.alignment = PP_ALIGN.RIGHT
    second_paragraph.space_after = Pt(3)

    table_shape = slide.shapes.add_table(2, 2, Inches(0.5), Inches(2.3), Inches(6.0), Inches(2.0))
    table_shape.name = "Formatted table"
    table = table_shape.table
    merged = table.cell(0, 0)
    merged.merge(table.cell(0, 1))
    merged.margin_left = Inches(0.09)
    merged.margin_right = Inches(0.11)
    merged.vertical_anchor = MSO_ANCHOR.BOTTOM
    merged.fill.solid()
    merged.fill.fore_color.rgb = RGBColor(0xEE, 0xDD, 0xCC)
    merged_p = merged.text_frame.paragraphs[0]
    merged_p.clear()
    merged_a = merged_p.add_run()
    merged_a.text = "Gross "
    merged_a.font.bold = True
    merged_b = merged_p.add_run()
    merged_b.text = "margin"
    merged_b.font.italic = True
    table.cell(1, 0).text = "North"
    table.cell(1, 1).text = "South"

    outer = slide.shapes.add_group_shape()
    outer.name = "Outer group"
    inner = outer.shapes.add_group_shape()
    inner.name = "Inner group"
    nested = inner.shapes.add_textbox(Inches(6.2), Inches(0.6), Inches(2.5), Inches(1.0))
    nested.name = "Nested text"
    nested.text_frame.paragraphs[0].text = "Nested group text"
    nested.rotation = 7

    presentation.slides.add_slide(presentation.slide_layouts[6])
    second_slide = presentation.slides[1]
    footer = second_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    footer.name = "Slide two"
    footer.text_frame.paragraphs[0].text = "Shared phrase"

    presentation.save(path)


def translated_values(plan, prefix: str = "TR") -> dict[str, dict[str, str]]:
    """Return deterministic translations without using a live provider."""

    return {
        unit.id: {span.id: f"{prefix}<{span.source}>" for span in unit.spans if span.translatable}
        for unit in plan.units
    }
