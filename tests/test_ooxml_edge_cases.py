from __future__ import annotations

import os
import stat
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from test_ooxml_helpers import create_complex_deck, translated_values

from pptrans.application.deck import write_translated_deck
from pptrans.domain import (
    InvalidPresentationError,
    LocatorResolutionError,
    PackageLimits,
    PatchSet,
    PatchValidationError,
    TextContainer,
    TranslatedSpan,
    VerificationError,
)
from pptrans.ooxml import apply_patch_set, build_patch_set, inspect_deck, verify_output
from pptrans.ooxml.locate import paragraph_for_locator, resolve_shape
from pptrans.ooxml.package import open_package, rewrite_package
from pptrans.ooxml.xml import (
    NS,
    assign_text,
    parse_xml,
    serialize_xml,
    structural_fingerprint,
    validate_text,
)


def _valid_plan(tmp_path: Path):
    source = tmp_path / "complex.pptx"
    create_complex_deck(source)
    return source, inspect_deck(source, source_lang="en", target_lang="fr")


def test_atomic_writer_rejects_unsafe_output_paths(tmp_path: Path) -> None:
    source, plan = _valid_plan(tmp_path)
    translations = translated_values(plan)

    with pytest.raises(PatchValidationError, match="must not overwrite"):
        write_translated_deck(plan, translations, source)
    with pytest.raises(PatchValidationError, match="extension"):
        write_translated_deck(plan, translations, tmp_path / "output.bin")

    existing = tmp_path / "existing.pptx"
    existing.write_bytes(b"existing")
    with pytest.raises(FileExistsError):
        write_translated_deck(plan, translations, existing)

    directory = tmp_path / "directory.pptx"
    directory.mkdir()
    with pytest.raises(IsADirectoryError):
        write_translated_deck(plan, translations, directory, overwrite=True)


def test_patch_builder_rejects_bad_provider_contracts(tmp_path: Path) -> None:
    _, plan = _valid_plan(tmp_path)
    translations = translated_values(plan)
    first = plan.units[0]

    with pytest.raises(PatchValidationError, match="Unsupported plan schema"):
        build_patch_set(replace(plan, schema_version="bad"), translations)
    with pytest.raises(PatchValidationError, match="Unknown translation unit"):
        build_patch_set(plan, {"unknown": {}})
    with pytest.raises(PatchValidationError, match="Missing translation unit"):
        build_patch_set(plan, {})

    with_extra = translated_values(plan)
    with_extra[first.id]["invented"] = "bad"
    with pytest.raises(PatchValidationError, match="unexpected invented"):
        build_patch_set(plan, with_extra)

    bad_character = translated_values(plan)
    bad_character[first.id][first.translatable_span_ids[0]] = "bad\x00text"
    with pytest.raises(PatchValidationError, match="forbidden by XML"):
        build_patch_set(plan, bad_character)

    duplicate = tuple(
        TranslatedSpan(span_id=span_id, text="ok") for span_id in first.translatable_span_ids
    )
    duplicate += (duplicate[0],)
    as_sequences = {
        unit.id: tuple(
            TranslatedSpan(span_id=span.id, text=span.source)
            for span in unit.spans
            if span.translatable
        )
        for unit in plan.units
    }
    as_sequences[first.id] = duplicate
    with pytest.raises(PatchValidationError, match="Duplicate translated span"):
        build_patch_set(plan, as_sequences)


def test_apply_rejects_tampered_patch_metadata(tmp_path: Path) -> None:
    source, plan = _valid_plan(tmp_path)
    patches = build_patch_set(plan, translated_values(plan))
    first = next(patch for patch in patches.patches if len(patch.translations) > 1)

    with pytest.raises(PatchValidationError, match="Unsupported patch schema"):
        apply_patch_set(
            source, tmp_path / "bad-schema.pptx", replace(patches, schema_version="bad")
        )

    duplicate = replace(patches, patches=(first, first))
    with pytest.raises(PatchValidationError, match="same paragraph"):
        apply_patch_set(source, tmp_path / "duplicate.pptx", duplicate)

    reordered = replace(first, translations=tuple(reversed(first.translations)))
    with pytest.raises(PatchValidationError, match="span order"):
        apply_patch_set(
            source,
            tmp_path / "reordered.pptx",
            PatchSet(patches.schema_version, patches.input_sha256, (reordered,)),
        )

    missing_locator = replace(
        first,
        locator=replace(first.locator, slide_part="ppt/slides/missing.xml"),
    )
    with pytest.raises(PatchValidationError, match="missing slide part"):
        apply_patch_set(
            source,
            tmp_path / "missing.pptx",
            PatchSet(patches.schema_version, patches.input_sha256, (missing_locator,)),
        )


def test_low_level_patch_and_rewrite_reject_source_aliases_without_data_loss(
    tmp_path: Path,
) -> None:
    source, plan = _valid_plan(tmp_path)
    original = source.read_bytes()
    patches = build_patch_set(plan, translated_values(plan))

    with pytest.raises(InvalidPresentationError, match="different files"):
        apply_patch_set(source, source, patches)
    assert source.read_bytes() == original

    with pytest.raises(InvalidPresentationError, match="different files"):
        rewrite_package(source, source, {})
    assert source.read_bytes() == original

    alias = tmp_path / "source-hard-link.pptx"
    try:
        os.link(source, alias)
    except OSError as exc:
        pytest.skip(f"hard links unavailable on this filesystem: {exc}")
    with pytest.raises(InvalidPresentationError, match="different files"):
        apply_patch_set(source, alias, patches)
    assert source.read_bytes() == original
    assert alias.read_bytes() == original


def test_low_level_explicit_overwrite_rejects_destination_symlink(tmp_path: Path) -> None:
    source, plan = _valid_plan(tmp_path)
    patches = build_patch_set(plan, translated_values(plan))
    target = tmp_path / "unrelated.pptx"
    sentinel = b"must remain untouched"
    target.write_bytes(sentinel)
    destination = tmp_path / "output-link.pptx"
    try:
        destination.symlink_to(target)
    except OSError as exc:
        pytest.skip(f"symbolic links unavailable on this filesystem: {exc}")

    with pytest.raises(InvalidPresentationError, match="symbolic link"):
        rewrite_package(source, destination, {}, overwrite=True)
    with pytest.raises(InvalidPresentationError, match="symbolic link"):
        apply_patch_set(source, destination, patches, overwrite=True)
    assert target.read_bytes() == sentinel


def test_low_level_patch_and_rewrite_do_not_clobber_by_default(tmp_path: Path) -> None:
    source, plan = _valid_plan(tmp_path)
    patches = build_patch_set(plan, translated_values(plan))
    destination = tmp_path / "existing.pptx"
    sentinel = b"preserve existing destination"
    destination.write_bytes(sentinel)

    with pytest.raises(FileExistsError, match="already exists"):
        apply_patch_set(source, destination, patches)
    assert destination.read_bytes() == sentinel

    with pytest.raises(FileExistsError, match="already exists"):
        rewrite_package(source, destination, {})
    assert destination.read_bytes() == sentinel


def test_low_level_patch_and_rewrite_support_explicit_overwrite(tmp_path: Path) -> None:
    source, plan = _valid_plan(tmp_path)
    patches = build_patch_set(plan, translated_values(plan))
    rewritten = tmp_path / "rewritten.pptx"
    rewritten.write_bytes(b"replace me")

    rewrite_package(source, rewritten, {}, overwrite=True)
    assert rewritten.read_bytes() == source.read_bytes()

    apply_patch_set(source, rewritten, patches, overwrite=True)
    report = verify_output(source, rewritten, patches)
    assert report.verified_patches == len(patches.patches)


def test_inspection_rejects_duplicate_presentation_relationship_ids(
    tmp_path: Path,
) -> None:
    source, _ = _valid_plan(tmp_path)
    malformed = tmp_path / "duplicate-relationship-id.pptx"
    relationships_name = "ppt/_rels/presentation.xml.rels"
    with zipfile.ZipFile(source) as archive:
        root = parse_xml(
            archive.read(relationships_name),
            part_name=relationships_name,
        )
    relationships = root.findall("rel:Relationship", NS)
    assert len(relationships) >= 2
    relationships[1].set("Id", relationships[0].get("Id"))
    rewrite_package(
        source,
        malformed,
        {relationships_name: serialize_xml(root)},
    )

    with pytest.raises(InvalidPresentationError, match="repeat relationship ID"):
        inspect_deck(malformed, source_lang="en", target_lang="fr")


def test_inspection_rejects_duplicate_slide_relationship_references(
    tmp_path: Path,
) -> None:
    source, _ = _valid_plan(tmp_path)
    malformed = tmp_path / "duplicate-slide-relationship.pptx"
    presentation_name = "ppt/presentation.xml"
    with zipfile.ZipFile(source) as archive:
        root = parse_xml(
            archive.read(presentation_name),
            part_name=presentation_name,
        )
    slide_ids = root.findall("./p:sldIdLst/p:sldId", NS)
    assert len(slide_ids) >= 2
    slide_ids[1].set(f"{{{NS['r']}}}id", slide_ids[0].get(f"{{{NS['r']}}}id"))
    rewrite_package(
        source,
        malformed,
        {presentation_name: serialize_xml(root)},
    )

    with pytest.raises(InvalidPresentationError, match="repeats relationship reference"):
        inspect_deck(malformed, source_lang="en", target_lang="fr")


def test_inspection_rejects_duplicate_slide_part_references(tmp_path: Path) -> None:
    source, _ = _valid_plan(tmp_path)
    malformed = tmp_path / "duplicate-slide-part.pptx"
    relationships_name = "ppt/_rels/presentation.xml.rels"
    with zipfile.ZipFile(source) as archive:
        root = parse_xml(
            archive.read(relationships_name),
            part_name=relationships_name,
        )
    slide_relationships = [
        relationship
        for relationship in root.findall("rel:Relationship", NS)
        if relationship.get("Type", "").rsplit("/", 1)[-1] == "slide"
    ]
    assert len(slide_relationships) >= 2
    slide_relationships[1].set("Target", slide_relationships[0].get("Target"))
    rewrite_package(
        source,
        malformed,
        {relationships_name: serialize_xml(root)},
    )

    with pytest.raises(InvalidPresentationError, match="repeats slide part reference"):
        inspect_deck(malformed, source_lang="en", target_lang="fr")


def test_package_limits_missing_members_and_invalid_zip(tmp_path: Path) -> None:
    missing = tmp_path / "missing.pptx"
    with pytest.raises(InvalidPresentationError, match="does not exist"), open_package(missing):
        pass

    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"not a zip")
    with pytest.raises(InvalidPresentationError, match="readable PPTX"), open_package(broken):
        pass

    incomplete = tmp_path / "incomplete.pptx"
    with zipfile.ZipFile(incomplete, "w") as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
    with (
        pytest.raises(InvalidPresentationError, match="missing required"),
        open_package(incomplete),
    ):
        pass

    source, _ = _valid_plan(tmp_path)
    with (
        pytest.raises(InvalidPresentationError, match="members; limit"),
        open_package(source, PackageLimits(max_members=1)),
    ):
        pass
    with (
        pytest.raises(InvalidPresentationError, match="size limit"),
        open_package(source, PackageLimits(max_member_bytes=1)),
    ):
        pass
    with (
        pytest.raises(InvalidPresentationError, match="XML safety limit"),
        open_package(source, PackageLimits(max_xml_bytes=1)),
    ):
        pass
    with (
        pytest.raises(InvalidPresentationError, match="uncompressed size"),
        open_package(
            source,
            PackageLimits(max_member_bytes=10**9, max_total_bytes=1),
        ),
    ):
        pass


def test_default_package_limits_bound_xml_and_archive_expansion() -> None:
    limits = PackageLimits()

    assert limits.max_xml_bytes == 32 * 1024 * 1024
    assert limits.max_member_bytes == 256 * 1024 * 1024
    assert limits.max_total_bytes == 1024 * 1024 * 1024
    assert limits.max_compression_ratio == 500
    assert limits.max_slides == 500
    assert limits.max_xml_elements_per_part == 250_000
    assert limits.max_translation_units == 10_000
    assert limits.max_text_spans == 50_000
    assert limits.max_source_characters == 5_000_000
    assert limits.max_translatable_spans_per_unit == 10_000
    assert limits.max_source_characters_per_span == 100_000
    assert limits.max_diagnostics == 10_000


@pytest.mark.parametrize(
    ("limits", "message"),
    [
        (PackageLimits(max_slides=0), "slides; limit"),
        (PackageLimits(max_xml_elements_per_part=1), "element-count"),
        (PackageLimits(max_translation_units=0), "translation-unit"),
        (PackageLimits(max_text_spans=0), "text-span"),
        (PackageLimits(max_source_characters=0), "source-character"),
        (
            PackageLimits(max_translatable_spans_per_unit=0),
            "translatable-span",
        ),
        (
            PackageLimits(max_source_characters_per_span=0),
            "per-span character",
        ),
    ],
    ids=[
        "slides",
        "elements",
        "units",
        "spans",
        "characters",
        "unit-spans",
        "span-characters",
    ],
)
def test_inspection_enforces_incremental_semantic_limits(
    tmp_path: Path,
    limits: PackageLimits,
    message: str,
) -> None:
    source, _ = _valid_plan(tmp_path)

    with pytest.raises(InvalidPresentationError, match=message):
        inspect_deck(source, source_lang="en", target_lang="fr", limits=limits)


def test_package_rejects_drive_paths_and_symbolic_links(tmp_path: Path) -> None:
    source, _ = _valid_plan(tmp_path)
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr("C:/outside.xml", b"unsafe")
    with (
        pytest.raises(InvalidPresentationError, match="Unsafe package member"),
        open_package(source),
    ):
        pass

    source.unlink()
    create_complex_deck(source)
    link = zipfile.ZipInfo("ppt/fake-link")
    link.external_attr = (stat.S_IFLNK | 0o777) << 16
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr(link, b"target")
    with pytest.raises(InvalidPresentationError, match="Symbolic-link"), open_package(source):
        pass

    source.unlink()
    create_complex_deck(source)
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr("_xmlsignatures/sig1.xml", b"<signature/>")
    with pytest.raises(InvalidPresentationError, match="Digitally signed"), open_package(source):
        pass


def test_locator_failures_are_explicit(tmp_path: Path) -> None:
    source, plan = _valid_plan(tmp_path)
    with zipfile.ZipFile(source) as archive:
        root = parse_xml(archive.read(plan.slide_parts[0]), part_name=plan.slide_parts[0])

    with pytest.raises(LocatorResolutionError, match="empty"):
        resolve_shape(root, ())
    with pytest.raises(LocatorResolutionError, match="Could not resolve"):
        resolve_shape(root, (999_999,))

    table = next(unit for unit in plan.units if unit.locator.container is TextContainer.TABLE_CELL)
    with pytest.raises(LocatorResolutionError, match="missing row and column"):
        paragraph_for_locator(root, replace(table.locator, cell=None))
    with pytest.raises(LocatorResolutionError, match="row 999"):
        paragraph_for_locator(root, replace(table.locator, cell=(999, 0)))

    shape = next(unit for unit in plan.units if unit.locator.container is TextContainer.SHAPE)
    with pytest.raises(LocatorResolutionError, match="Paragraph 999"):
        paragraph_for_locator(root, replace(shape.locator, paragraph_index=999))


def test_inspection_reports_unsupported_graphic_frames(tmp_path: Path) -> None:
    source = tmp_path / "chart.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Series", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(5),
        Inches(3),
        data,
    )
    presentation.save(source)

    plan = inspect_deck(source, source_lang="en", target_lang="fr")
    assert [warning.code for warning in plan.warnings] == ["unsupported_graphic_frame"]
    with pytest.raises(ValueError, match="non-empty"):
        inspect_deck(source, source_lang="", target_lang="fr")

    with pytest.raises(InvalidPresentationError, match="diagnostic safety limit"):
        inspect_deck(
            source,
            source_lang="en",
            target_lang="fr",
            limits=PackageLimits(max_diagnostics=0),
        )


def test_rejects_duplicate_sibling_shape_ids(tmp_path: Path) -> None:
    source, plan = _valid_plan(tmp_path)
    malformed = tmp_path / "duplicate-shape-id.pptx"
    slide_part = plan.slide_parts[0]
    with zipfile.ZipFile(source) as archive:
        root = parse_xml(archive.read(slide_part), part_name=slide_part)
    identifiers = root.xpath("./p:cSld/p:spTree/*/*/p:cNvPr", namespaces=NS)
    assert len(identifiers) >= 2
    identifiers[1].set("id", identifiers[0].get("id"))
    rewrite_package(source, malformed, {slide_part: serialize_xml(root)})

    with pytest.raises(InvalidPresentationError, match="repeat non-visual shape ID"):
        inspect_deck(malformed, source_lang="en", target_lang="fr")


def test_xml_validation_and_whitespace_semantics() -> None:
    with pytest.raises(InvalidPresentationError, match="DTD"):
        parse_xml(b"<!DOCTYPE x><x/>", part_name="unsafe.xml")
    with pytest.raises(InvalidPresentationError, match="Invalid XML"):
        parse_xml(b"<broken>", part_name="broken.xml")
    with pytest.raises(PatchValidationError, match="forbidden"):
        validate_text("bad\x01text")

    node = parse_xml(
        b'<a:t xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">x</a:t>',
        part_name="text.xml",
    )
    assign_text(node, " translated ")
    assert node.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"


def test_long_xml_declaration_is_detected_and_preserved() -> None:
    declaration = b'<?xml version="1.0"' + b" " * 5_000 + b"?>"
    original = declaration + b"<root/>"
    root = parse_xml(original, part_name="long-declaration.xml")

    serialized = serialize_xml(root, original=original)

    assert serialized.startswith(declaration)
    assert structural_fingerprint(
        original, part_name="long-declaration.xml"
    ) == structural_fingerprint(serialized, part_name="long-declaration.xml")


@pytest.mark.parametrize(
    "declaration",
    [
        b'<?xml version="1.0" encoding="UTF-8"?>',
        b'<?xml version="1.0" encoding="UTF-8" standalone="no"?>',
    ],
    ids=["no-standalone", "standalone-no"],
)
def test_changed_slide_preserves_full_xml_document_and_declaration(
    tmp_path: Path,
    declaration: bytes,
) -> None:
    source, plan = _valid_plan(tmp_path)
    decorated_source = tmp_path / "decorated.pptx"
    output = tmp_path / "decorated-translated.pptx"
    slide_part = plan.slide_parts[0]
    with zipfile.ZipFile(source) as archive:
        original = archive.read(slide_part)
    body = original.split(b"?>", 1)[1] if original.startswith(b"<?xml") else original
    decorated = declaration + b"<?before x?><!--pre-root-->" + body + b"<!--post-root--><?after y?>"
    rewrite_package(source, decorated_source, {slide_part: decorated})
    decorated_plan = inspect_deck(decorated_source, source_lang="en", target_lang="fr")
    patches = build_patch_set(decorated_plan, translated_values(decorated_plan))

    apply_patch_set(decorated_source, output, patches)

    with zipfile.ZipFile(output) as archive:
        changed = archive.read(slide_part)
    for marker in (b"<?before x?>", b"<!--pre-root-->", b"<!--post-root-->", b"<?after y?>"):
        assert marker in changed
    declaration_out = changed.split(b"?>", 1)[0]
    assert (b"standalone" in declaration_out) == (b"standalone" in declaration)
    assert structural_fingerprint(decorated, part_name=slide_part) == structural_fingerprint(
        changed, part_name=slide_part
    )
    assert structural_fingerprint(
        changed.replace(b"<?before x?>", b"<?before changed?>"),
        part_name=slide_part,
    ) != structural_fingerprint(changed, part_name=slide_part)


@pytest.mark.parametrize(
    "document",
    [
        "<?xml version='1.0'?><!DOCTYPE root [<!ENTITY value 'unsafe'>]><root>&value;</root>",
        "<?xml version='1.0'?><!DOCTYPE root SYSTEM 'https://untrusted.invalid/dtd'><root/>",
    ],
    ids=["internal-entity", "external-system"],
)
@pytest.mark.parametrize("encoding", ["utf-16", "utf-32"])
def test_xml_validation_rejects_dtds_independently_of_encoding(
    document: str,
    encoding: str,
) -> None:
    with pytest.raises(InvalidPresentationError, match="DTD"):
        parse_xml(document.encode(encoding), part_name=f"unsafe-{encoding}.xml")


def test_verifier_rejects_inventory_and_structure_changes(tmp_path: Path) -> None:
    source, plan = _valid_plan(tmp_path)
    staged = tmp_path / "staged.pptx"
    structural = tmp_path / "structural.pptx"
    inventory = tmp_path / "inventory.pptx"
    patches = build_patch_set(plan, translated_values(plan))
    apply_patch_set(source, staged, patches)

    slide_part = plan.slide_parts[0]
    with zipfile.ZipFile(staged) as archive:
        root = parse_xml(archive.read(slide_part), part_name=slide_part)
    shape_properties = root.find(".//p:spPr", NS)
    assert shape_properties is not None
    shape_properties.set("unexpected", "true")
    rewrite_package(staged, structural, {slide_part: serialize_xml(root)})
    with pytest.raises(VerificationError, match="Formatting structure changed"):
        verify_output(source, structural, patches)

    inventory.write_bytes(staged.read_bytes())
    with zipfile.ZipFile(inventory, "a") as archive:
        archive.writestr("extra.xml", b"<extra/>")
    with pytest.raises(VerificationError, match="member names or ordering"):
        verify_output(source, inventory, patches)
