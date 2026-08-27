from __future__ import annotations

import hashlib
import os
import subprocess
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
from PIL import Image
from pptx import Presentation

from pptrans.adapters.renderers import (
    LibreOfficeRenderer,
    RenderError,
    RendererUnavailable,
    RenderRequest,
    UnsafeRenderInput,
)


class FakeRasterizer:
    def __init__(self, page_count: int = 1) -> None:
        self._page_count = page_count

    @staticmethod
    def available() -> bool:
        return True

    def rasterize(
        self,
        pdf_path: Path,
        output_dir: Path,
        *,
        dpi: int,
        max_pages: int,
        max_total_pixels: int,
    ) -> tuple[Path, ...]:
        assert pdf_path.read_bytes() == b"%PDF-fake"
        assert dpi == 144
        assert max_pages >= 1
        assert max_total_pixels >= 800
        image_paths = []
        for page_number in range(1, self._page_count + 1):
            image_path = output_dir / f"page-{page_number:04d}.png"
            Image.new("RGB", (40, 20), "white").save(image_path, "PNG")
            image_paths.append(image_path)
        return tuple(image_paths)


def _renderer(
    tmp_path: Path,
    calls: list[tuple[list[str], dict[str, Any]]],
    *,
    page_count: int = 1,
    slide_count: int | None = 1,
    linker: Callable[[Path, Path], None] = os.link,
) -> LibreOfficeRenderer:
    executable = tmp_path / "soffice.exe"
    executable.write_bytes(b"fake executable")

    def fake_runner(args: list[str], **kwargs: Any) -> subprocess.CompletedProcess[str]:
        calls.append((args, kwargs))
        conversion_dir = Path(args[args.index("--outdir") + 1])
        source = Path(args[-1])
        (conversion_dir / f"{source.stem}.pdf").write_bytes(b"%PDF-fake")
        return subprocess.CompletedProcess(args, 0, "", "")

    if slide_count is None:
        return LibreOfficeRenderer(
            executable=executable,
            rasterizer=FakeRasterizer(page_count),
            runner=fake_runner,
            environment={"PATH": "safe", "OPENAI_API_KEY": "must-not-leak"},
            linker=linker,
        )
    return LibreOfficeRenderer(
        executable=executable,
        rasterizer=FakeRasterizer(page_count),
        runner=fake_runner,
        environment={"PATH": "safe", "OPENAI_API_KEY": "must-not-leak"},
        slide_counter=lambda _source: slide_count,
        linker=linker,
    )


def _create_pptx(path: Path, *, slide_count: int) -> None:
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for _index in range(slide_count):
        presentation.slides.add_slide(blank)
    presentation.save(str(path))


def test_libreoffice_renderer_uses_argument_list_and_isolated_profile(tmp_path: Path) -> None:
    calls: list[tuple[list[str], dict[str, Any]]] = []
    renderer = _renderer(tmp_path, calls)
    source = tmp_path / "deck;do-not-run.pptx"
    source.write_bytes(b"fake pptx")
    output_dir = tmp_path / "renders"

    result = renderer.render(RenderRequest(source, output_dir))

    assert result.backend == "libreoffice"
    assert len(result.slides) == 1
    assert result.slides[0].image_path.is_file()
    assert result.slides[0].width == 40
    args, kwargs = calls[0]
    rendered_source = Path(args[-1])
    assert rendered_source.name == source.name
    assert rendered_source != source.resolve()
    assert rendered_source.parent.parent == output_dir.resolve()
    assert kwargs["shell"] is False
    assert kwargs["check"] is False
    assert kwargs["cwd"]
    assert any(argument.startswith("-env:UserInstallation=file:") for argument in args)
    export_filter = args[args.index("--convert-to") + 1]
    assert export_filter.startswith("pdf:impress_pdf_Export:")
    assert '"ExportHiddenSlides"' in export_filter
    assert '"value":"true"' in export_filter
    assert "OPENAI_API_KEY" not in kwargs["env"]
    assert not (tmp_path / "do-not-run").exists()


def test_renderer_uses_one_private_source_snapshot(tmp_path: Path) -> None:
    executable = tmp_path / "soffice.exe"
    executable.write_bytes(b"fake executable")
    source = tmp_path / "deck.pptx"
    original = b"source version used for rendering"
    source.write_bytes(original)

    def fake_runner(args: list[str], **_kwargs: Any) -> subprocess.CompletedProcess[str]:
        staged_source = Path(args[-1])
        assert staged_source != source.resolve()
        assert staged_source.read_bytes() == original
        source.write_bytes(b"concurrent replacement")
        conversion_dir = Path(args[args.index("--outdir") + 1])
        (conversion_dir / f"{staged_source.stem}.pdf").write_bytes(b"%PDF-fake")
        return subprocess.CompletedProcess(args, 0, "", "")

    renderer = LibreOfficeRenderer(
        executable=executable,
        rasterizer=FakeRasterizer(),
        runner=fake_runner,
        slide_counter=lambda staged: 1 if staged.read_bytes() == original else 0,
    )
    result = renderer.render(RenderRequest(source, tmp_path / "renders"))

    assert result.source_sha256 == hashlib.sha256(original).hexdigest()
    assert source.read_bytes() == b"concurrent replacement"


def test_renderer_reports_missing_executable(tmp_path: Path) -> None:
    renderer = LibreOfficeRenderer(executable=tmp_path / "missing", rasterizer=FakeRasterizer())
    availability = renderer.availability()
    assert availability.available is False
    source = tmp_path / "deck.pptx"
    source.write_bytes(b"pptx")
    with pytest.raises(RendererUnavailable):
        renderer.render(RenderRequest(source, tmp_path / "out"))


def test_renderer_rejects_legacy_and_non_presentation_inputs(tmp_path: Path) -> None:
    calls: list[tuple[list[str], dict[str, Any]]] = []
    renderer = _renderer(tmp_path, calls)
    legacy = tmp_path / "legacy.ppt"
    legacy.write_bytes(b"legacy")
    with pytest.raises(UnsafeRenderInput, match=r"only \.pptx"):
        renderer.render(RenderRequest(legacy, tmp_path / "out"))
    text = tmp_path / "deck.txt"
    text.write_text("not a deck", encoding="utf-8")
    with pytest.raises(UnsafeRenderInput, match=r"only \.pptx"):
        renderer.render(RenderRequest(text, tmp_path / "out"))
    assert calls == []


def test_renderer_rejects_invalid_raster_output(tmp_path: Path) -> None:
    class EmptyRasterizer(FakeRasterizer):
        def rasterize(self, *args: Any, **kwargs: Any) -> tuple[Path, ...]:
            return ()

    executable = tmp_path / "soffice.exe"
    executable.write_bytes(b"fake")

    def fake_runner(args: list[str], **kwargs: Any) -> subprocess.CompletedProcess[str]:
        conversion_dir = Path(args[args.index("--outdir") + 1])
        source = Path(args[-1])
        (conversion_dir / f"{source.stem}.pdf").write_bytes(b"%PDF-fake")
        return subprocess.CompletedProcess(args, 0, "", "")

    source = tmp_path / "deck.pptx"
    source.write_bytes(b"pptx")
    renderer = LibreOfficeRenderer(
        executable=executable,
        rasterizer=EmptyRasterizer(),
        runner=fake_runner,
        slide_counter=lambda _source: 1,
    )
    with pytest.raises(RenderError, match="no slide images"):
        renderer.render(RenderRequest(source, tmp_path / "out"))


def test_renderer_rejects_dropped_raster_pages_before_publication(tmp_path: Path) -> None:
    calls: list[tuple[list[str], dict[str, Any]]] = []
    renderer = _renderer(tmp_path, calls, page_count=1, slide_count=None)
    source = tmp_path / "two-slides.pptx"
    _create_pptx(source, slide_count=2)
    output_dir = tmp_path / "renders"

    with pytest.raises(RenderError, match=r"does not match.*slide count"):
        renderer.render(RenderRequest(source, output_dir))

    assert list(output_dir.glob("*.png")) == []


def test_renderer_refuses_an_existing_render_without_modifying_it(tmp_path: Path) -> None:
    calls: list[tuple[list[str], dict[str, Any]]] = []
    renderer = _renderer(tmp_path, calls)
    source = tmp_path / "deck.pptx"
    source.write_bytes(b"synthetic pptx")
    output_dir = tmp_path / "renders"
    output_dir.mkdir()
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    collision = output_dir / f"deck-{digest[:12]}-slide-0001.png"
    collision.write_bytes(b"existing render")

    with pytest.raises(RenderError, match="refusing to overwrite"):
        renderer.render(RenderRequest(source, output_dir))

    assert collision.read_bytes() == b"existing render"
    assert list(output_dir.glob("*.png")) == [collision]


def test_renderer_rolls_back_a_mid_publication_failure(tmp_path: Path) -> None:
    calls: list[tuple[list[str], dict[str, Any]]] = []
    link_count = 0

    def fail_second_link(source: Path, destination: Path) -> None:
        nonlocal link_count
        link_count += 1
        os.link(source, destination)
        if link_count == 2:
            raise OSError("injected link failure")

    renderer = _renderer(
        tmp_path,
        calls,
        page_count=2,
        slide_count=2,
        linker=fail_second_link,
    )
    source = tmp_path / "deck.pptx"
    source.write_bytes(b"synthetic pptx")
    output_dir = tmp_path / "renders"

    with pytest.raises(RenderError, match=r"no-clobber.*no files published"):
        renderer.render(RenderRequest(source, output_dir))

    assert link_count == 2
    assert list(output_dir.glob("*.png")) == []


def test_renderer_publishes_every_page_after_real_pptx_count_check(tmp_path: Path) -> None:
    calls: list[tuple[list[str], dict[str, Any]]] = []
    executable = tmp_path / "soffice.exe"
    executable.write_bytes(b"fake executable")

    def fake_runner(args: list[str], **_kwargs: Any) -> subprocess.CompletedProcess[str]:
        calls.append((args, _kwargs))
        conversion_dir = Path(args[args.index("--outdir") + 1])
        source = Path(args[-1])
        (conversion_dir / f"{source.stem}.pdf").write_bytes(b"%PDF-fake")
        return subprocess.CompletedProcess(args, 0, "", "")

    source = tmp_path / "deck.pptx"
    _create_pptx(source, slide_count=2)
    renderer = LibreOfficeRenderer(
        executable=executable,
        rasterizer=FakeRasterizer(page_count=2),
        runner=fake_runner,
    )

    result = renderer.render(RenderRequest(source, tmp_path / "renders"))

    assert [slide.slide_number for slide in result.slides] == [1, 2]
    assert all(slide.image_path.is_file() for slide in result.slides)
    assert all(
        slide.sha256 == hashlib.sha256(slide.image_path.read_bytes()).hexdigest()
        for slide in result.slides
    )
