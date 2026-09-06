from __future__ import annotations

import errno
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from test_ooxml_helpers import create_complex_deck, translated_values

import pptrans.application.deck as deck_service
from pptrans.application.deck import write_translated_deck
from pptrans.domain import PatchValidationError, SourceChangedError, TextContainer
from pptrans.ooxml import inspect_deck, package_member_hashes
from pptrans.ooxml.package import file_sha256
from pptrans.ooxml.xml import structural_fingerprint


def _slide_bytes(path: Path, part: str) -> bytes:
    with zipfile.ZipFile(path) as archive:
        return archive.read(part)


def test_inspection_builds_stable_run_table_and_group_units(tmp_path: Path) -> None:
    source = tmp_path / "complex.pptx"
    create_complex_deck(source)

    first = inspect_deck(source, source_lang="en", target_lang="fr")
    second = inspect_deck(source, source_lang="en", target_lang="fr")

    assert first.input_sha256 == file_sha256(source)
    assert first.slide_parts == ("ppt/slides/slide1.xml", "ppt/slides/slide2.xml")
    assert [unit.id for unit in first.units] == [unit.id for unit in second.units]
    assert len(first.units) >= 7

    mixed = next(unit for unit in first.units if unit.source_text.startswith("Revenue "))
    assert [span.source for span in mixed.spans] == [
        "Revenue ",
        "grew 20%",
        "2026-08-28",
        "Year over year",
    ]
    assert mixed.translatable_span_ids == ("s0", "s1", "s3")

    table_units = [
        unit for unit in first.units if unit.locator.container is TextContainer.TABLE_CELL
    ]
    assert {unit.locator.cell for unit in table_units} == {(0, 0), (1, 0), (1, 1)}
    assert any(len(unit.locator.shape_id_path) == 3 for unit in first.units)
    assert first.units[-1].locator.slide_index == 2


def test_atomic_translation_changes_only_allowed_text(tmp_path: Path) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "complex_translated.pptx"
    create_complex_deck(source)
    source_before = source.read_bytes()
    plan = inspect_deck(source, source_lang="en", target_lang="fr")

    result = write_translated_deck(plan, translated_values(plan), output)

    assert result.output_path == output.resolve()
    assert source.read_bytes() == source_before
    assert result.report.verified_patches == len(plan.units)
    assert result.report.verified_spans == sum(
        len(unit.translatable_span_ids) for unit in plan.units
    )
    assert set(result.report.changed_parts).issubset(set(plan.slide_parts))

    source_hashes = package_member_hashes(source)
    output_hashes = package_member_hashes(output)
    for member, digest in source_hashes.items():
        if member not in plan.slide_parts:
            assert output_hashes[member] == digest
    for slide_part in plan.slide_parts:
        assert structural_fingerprint(
            _slide_bytes(source, slide_part), part_name=slide_part
        ) == structural_fingerprint(_slide_bytes(output, slide_part), part_name=slide_part)

    # A second independent consumer must be able to open the package.
    reopened = Presentation(output)
    assert len(reopened.slides) == 2
    translated_shape = next(
        shape for shape in reopened.slides[0].shapes if shape.name == "Mixed runs"
    )
    runs = translated_shape.text_frame.paragraphs[0].runs
    assert runs[0].text == "TR<Revenue >"
    assert runs[0].font.bold is True
    assert runs[0].font.size.pt == 24
    assert str(runs[0].font.color.rgb) == "123456"
    assert runs[0].hyperlink.address == "https://example.com/revenue"
    assert runs[1].text == "TR<grew 20%>"
    assert runs[1].font.italic is True
    assert runs[1].font.size.pt == 18


def test_identity_translation_keeps_every_package_part_byte_identical(tmp_path: Path) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "identity.pptx"
    create_complex_deck(source)
    plan = inspect_deck(source, source_lang="en", target_lang="en")
    identity = {
        unit.id: {span.id: span.source for span in unit.spans if span.translatable}
        for unit in plan.units
    }

    result = write_translated_deck(plan, identity, output)

    assert result.report.changed_parts == ()
    assert package_member_hashes(output) == package_member_hashes(source)


def test_malformed_span_map_fails_before_creating_output(tmp_path: Path) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "bad.pptx"
    create_complex_deck(source)
    plan = inspect_deck(source, source_lang="en", target_lang="fr")
    translations = translated_values(plan)
    first_unit = plan.units[0]
    translations[first_unit.id].pop(first_unit.translatable_span_ids[0])

    with pytest.raises(PatchValidationError, match="invalid spans"):
        write_translated_deck(plan, translations, output)

    assert not output.exists()
    assert not tuple(tmp_path.glob(".bad.pptrans-*.pptx"))


def test_source_hash_guard_rejects_post_inspection_changes(tmp_path: Path) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "out.pptx"
    create_complex_deck(source)
    plan = inspect_deck(source, source_lang="en", target_lang="fr")
    translations = translated_values(plan)

    changed = Presentation(source)
    changed.slides.add_slide(changed.slide_layouts[6])
    changed.save(source)

    with pytest.raises(SourceChangedError, match="changed after inspection"):
        write_translated_deck(plan, translations, output)
    assert not output.exists()


def test_verification_failure_preserves_existing_destination(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "existing.pptx"
    create_complex_deck(source)
    output.write_bytes(b"keep me")
    plan = inspect_deck(source, source_lang="en", target_lang="fr")

    def fail_verification(*args, **kwargs):
        raise RuntimeError("injected verifier failure")

    monkeypatch.setattr(deck_service, "verify_output", fail_verification)
    with pytest.raises(RuntimeError, match="injected verifier failure"):
        write_translated_deck(plan, translated_values(plan), output, overwrite=True)

    assert output.read_bytes() == b"keep me"
    assert not tuple(tmp_path.glob(".existing.pptrans-*.pptx"))


def test_no_overwrite_publish_rejects_a_destination_created_during_the_run(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "raced.pptx"
    create_complex_deck(source)
    plan = inspect_deck(source, source_lang="en", target_lang="fr")
    original_verify = deck_service.verify_output

    def create_racing_output(*args, **kwargs):
        report = original_verify(*args, **kwargs)
        output.write_bytes(b"created by another process")
        return report

    monkeypatch.setattr(deck_service, "verify_output", create_racing_output)

    with pytest.raises(FileExistsError, match="appeared before publish"):
        write_translated_deck(plan, translated_values(plan), output)

    assert output.read_bytes() == b"created by another process"
    assert not tuple(tmp_path.glob(".raced.pptrans-*.pptx"))


def test_no_overwrite_publish_fails_closed_when_atomic_link_is_unavailable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "unsupported-filesystem.pptx"
    create_complex_deck(source)
    plan = inspect_deck(source, source_lang="en", target_lang="fr")

    def reject_link(*_args, **_kwargs):
        raise OSError("hard links unavailable")

    monkeypatch.setattr(deck_service.os, "link", reject_link)

    with pytest.raises(PatchValidationError, match="atomically publish"):
        write_translated_deck(plan, translated_values(plan), output)

    assert not output.exists()
    assert not tuple(tmp_path.glob(".unsupported-filesystem.pptrans-*.pptx"))


def test_no_overwrite_publish_retries_transient_stage_cleanup(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "retry-cleanup.pptx"
    create_complex_deck(source)
    plan = inspect_deck(source, source_lang="en", target_lang="fr")
    original_unlink = deck_service._unlink_path
    stage_attempts = 0
    sleeps: list[float] = []

    def transient_unlink(path: Path) -> None:
        nonlocal stage_attempts
        if path.name.startswith(".retry-cleanup.pptrans-"):
            stage_attempts += 1
            if stage_attempts < 3:
                error = PermissionError(errno.EACCES, "injected sharing violation")
                error.winerror = 32
                raise error
        original_unlink(path)

    monkeypatch.setattr(deck_service, "_unlink_path", transient_unlink)
    monkeypatch.setattr(deck_service.time, "sleep", sleeps.append)

    result = write_translated_deck(plan, translated_values(plan), output)

    assert result.output_path == output
    assert stage_attempts == 3
    assert sleeps == [0.05, 0.1]
    assert output.is_file()
    assert not tuple(tmp_path.glob(".retry-cleanup.pptrans-*.pptx"))


def test_persistent_stage_cleanup_rolls_back_only_the_owned_final(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "persistent-cleanup.pptx"
    create_complex_deck(source)
    plan = inspect_deck(source, source_lang="en", target_lang="fr")
    original_unlink = deck_service._unlink_path
    stage_attempts = 0
    sleeps: list[float] = []

    def persistent_unlink(path: Path) -> None:
        nonlocal stage_attempts
        if path.name.startswith(".persistent-cleanup.pptrans-"):
            stage_attempts += 1
            raise PermissionError(errno.EACCES, "injected persistent sharing violation")
        original_unlink(path)

    monkeypatch.setattr(deck_service, "_unlink_path", persistent_unlink)
    monkeypatch.setattr(deck_service.time, "sleep", sleeps.append)

    with pytest.raises(
        PatchValidationError,
        match=r"owned final output was rolled back; private staging remains",
    ):
        write_translated_deck(plan, translated_values(plan), output)

    assert stage_attempts == len(deck_service._UNLINK_RETRY_DELAYS_SECONDS) + 1
    assert sleeps == list(deck_service._UNLINK_RETRY_DELAYS_SECONDS)
    assert not output.exists()
    assert len(tuple(tmp_path.glob(".persistent-cleanup.pptrans-*.pptx"))) == 1


def test_cleanup_rollback_preserves_a_foreign_final_replacement(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "foreign-replacement.pptx"
    sentinel = b"created by another process"
    create_complex_deck(source)
    plan = inspect_deck(source, source_lang="en", target_lang="fr")
    original_unlink = deck_service._unlink_path
    replaced = False

    def replace_final_then_fail(path: Path) -> None:
        nonlocal replaced
        if path.name.startswith(".foreign-replacement.pptrans-"):
            if not replaced:
                output.unlink()
                output.write_bytes(sentinel)
                replaced = True
            raise OSError(errno.EIO, "injected staging cleanup failure")
        original_unlink(path)

    monkeypatch.setattr(deck_service, "_unlink_path", replace_final_then_fail)

    with pytest.raises(
        PatchValidationError,
        match=r"foreign replacement at final output was preserved.*private staging remains",
    ):
        write_translated_deck(plan, translated_values(plan), output)

    assert replaced is True
    assert output.read_bytes() == sentinel
    assert len(tuple(tmp_path.glob(".foreign-replacement.pptrans-*.pptx"))) == 1


def test_cleanup_reports_owned_final_and_stage_when_both_unlinks_fail(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "double-residual.pptx"
    create_complex_deck(source)
    plan = inspect_deck(source, source_lang="en", target_lang="fr")

    def reject_owned_unlink(path: Path) -> None:
        if path == output or path.name.startswith(".double-residual.pptrans-"):
            raise PermissionError(errno.EACCES, "injected persistent sharing violation")
        path.unlink()

    monkeypatch.setattr(deck_service, "_unlink_path", reject_owned_unlink)
    monkeypatch.setattr(deck_service.time, "sleep", lambda _delay: None)

    with pytest.raises(
        PatchValidationError,
        match=r"owned final output remains.*private staging remains",
    ):
        write_translated_deck(plan, translated_values(plan), output)

    assert output.is_file()
    assert len(tuple(tmp_path.glob(".double-residual.pptrans-*.pptx"))) == 1


def test_prepublication_failure_survives_transient_stage_cleanup(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "prepublish-failure.pptx"
    create_complex_deck(source)
    plan = inspect_deck(source, source_lang="en", target_lang="fr")
    primary_error = RuntimeError("injected verifier failure")
    original_unlink = deck_service._unlink_path
    stage_attempts = 0

    def fail_verification(*_args: object, **_kwargs: object) -> None:
        raise primary_error

    def transient_unlink(path: Path) -> None:
        nonlocal stage_attempts
        if path.name.startswith(".prepublish-failure.pptrans-"):
            stage_attempts += 1
            if stage_attempts == 1:
                raise PermissionError(errno.EACCES, "injected sharing violation")
        original_unlink(path)

    monkeypatch.setattr(deck_service, "verify_output", fail_verification)
    monkeypatch.setattr(deck_service, "_unlink_path", transient_unlink)
    monkeypatch.setattr(deck_service.time, "sleep", lambda _delay: None)

    with pytest.raises(RuntimeError, match="injected verifier failure") as raised:
        write_translated_deck(plan, translated_values(plan), output)

    assert raised.value is primary_error
    assert stage_attempts == 2
    assert not output.exists()
    assert not tuple(tmp_path.glob(".prepublish-failure.pptrans-*.pptx"))


def test_prepublication_cleanup_error_retains_the_primary_failure_as_cause(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "prepublish-residual.pptx"
    create_complex_deck(source)
    plan = inspect_deck(source, source_lang="en", target_lang="fr")
    primary_error = RuntimeError("injected verifier failure")

    def fail_verification(*_args: object, **_kwargs: object) -> None:
        raise primary_error

    def reject_stage_unlink(path: Path) -> None:
        if path.name.startswith(".prepublish-residual.pptrans-"):
            raise PermissionError(errno.EACCES, "injected persistent sharing violation")
        path.unlink()

    monkeypatch.setattr(deck_service, "verify_output", fail_verification)
    monkeypatch.setattr(deck_service, "_unlink_path", reject_stage_unlink)
    monkeypatch.setattr(deck_service.time, "sleep", lambda _delay: None)

    with pytest.raises(PatchValidationError, match="staging cleanup was incomplete") as raised:
        write_translated_deck(plan, translated_values(plan), output)

    assert raised.value.__cause__ is primary_error
    assert not output.exists()
    assert len(tuple(tmp_path.glob(".prepublish-residual.pptrans-*.pptx"))) == 1


def test_overwrite_rechecks_for_a_symlink_created_during_the_run(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "complex.pptx"
    output = tmp_path / "raced-link.pptx"
    sentinel = tmp_path / "sentinel.bin"
    create_complex_deck(source)
    sentinel.write_bytes(b"preserve target")
    plan = inspect_deck(source, source_lang="en", target_lang="fr")
    original_verify = deck_service.verify_output

    def create_racing_symlink(*args, **kwargs):
        report = original_verify(*args, **kwargs)
        try:
            output.symlink_to(sentinel)
        except OSError as exc:
            pytest.skip(f"symbolic links unavailable on this filesystem: {exc}")
        return report

    monkeypatch.setattr(deck_service, "verify_output", create_racing_symlink)

    with pytest.raises(PatchValidationError, match="became a symbolic link"):
        write_translated_deck(plan, translated_values(plan), output, overwrite=True)

    assert sentinel.read_bytes() == b"preserve target"
    assert output.is_symlink()
    assert not tuple(tmp_path.glob(".raced-link.pptrans-*.pptx"))
